"""Genera presentacion-banco-plus.pptx a partir de presentacion-libreto.md y diagramas.py.

Mismo formato, tamaños (proporcionales) y colores que presentacion.html. El guion va en las notas del orador.
Uso: uv run --with python-pptx --with cairosvg python generar-pptx.py
"""
import re
from io import BytesIO
from pathlib import Path

import cairosvg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from diagramas import DIAGRAMAS

AQUI = Path(__file__).parent
md = (AQUI / "presentacion-libreto.md").read_text(encoding="utf-8")
FUENTE = "Arial"
AZUL, OSCURO, GRIS = "1168BD", "0B2A4A", "5B6470"
W, H = 13.333, 7.5
VH = H * 72 / 100  # 1vh del HTML, en puntos
MX = W * 0.04  # padding lateral del HTML (4vw)
EJES = [(1, "SLOs y Quality Gate"), (2, "Telemetría y FinOps"), (3, "IA y MCP"), (4, "Incidentes y cultura")]
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⬀-⯿️]")


def rgb(h):
    return RGBColor.from_string(h)


def caja(slide, x, y, w, h, relleno=None, borde=None, radio=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if radio:
        s.adjustments[0] = 0.5
    if relleno:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(relleno)
    else:
        s.fill.background()
    if borde:
        s.line.color.rgb = rgb(borde)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def runs(parrafo, texto, tam, color, negrita=False, italica=False):
    """Escribe texto con **negrita**, `código` e *itálica* como runs separados."""
    for trozo in re.split(r"(\*\*.+?\*\*|`.+?`|(?<!\w)\*[^*]+?\*(?!\w))", texto):
        if not trozo:
            continue
        r = parrafo.add_run()
        b, i, f = negrita, italica, FUENTE
        if trozo.startswith("**"):
            trozo, b = trozo[2:-2], True
        elif trozo.startswith("`"):
            trozo, f = trozo[1:-1], "Courier New"
        elif trozo.startswith("*"):
            trozo, i = trozo[1:-1], True
        r.text = trozo
        r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(tam), b, i, f
        r.font.color.rgb = rgb(color)


def texto(slide, x, y, w, h, contenido, tam, color, negrita=False, italica=False, alin=PP_ALIGN.LEFT,
          ancla=MSO_ANCHOR.TOP, margen=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    tf.margin_left = tf.margin_right = Inches(margen)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = alin
    runs(p, contenido, tam, color, negrita, italica)
    return tf


def bullets(bloque):
    items, actual = [], ""
    for linea in bloque.splitlines():
        if linea.startswith("- "):
            if actual:
                items.append(actual)
            actual = linea[2:]
        elif linea.strip():
            actual += " " + linea.strip()
    return items + ([actual] if actual else [])


def png(svg):
    svg = EMOJI.sub("", svg)
    return BytesIO(cairosvg.svg2png(bytestring=svg.encode(), output_width=2400))


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
vacia = prs.slide_layouts[6]

for parte in re.split(r"(?m)^## Diapositiva ", md)[1:]:
    num = int(parte.split(" ", 1)[0])
    tiempo = re.search(r"\(([\d,]+ min)\)", parte.splitlines()[0]).group(1)
    titulo = re.search(r"\*\*Título:\*\* (.+)", parte).group(1)
    pregunta = re.search(r"\*\*Pregunta:\*\* (.+)", parte).group(1)
    ejes = re.search(r"\*\*Ejes:\*\* (.+)", parte).group(1)
    puente = re.search(r"\*\*Puente a la siguiente:\*\* (.+)", parte).group(1)
    contenido = bullets(re.search(r"\*\*Contenido:\*\*\n(.*?)\n\n", parte, re.S).group(1))
    guion = re.search(r"\*\*Guion:\*\*\n(.*?)\n\n\*\*", parte, re.S).group(1).strip().strip('"')
    principal, usa = ejes.split("·")[0], (ejes.split("·")[1] if "·" in ejes else "")

    s = prs.slides.add_slide(vacia)
    caja(s, 0, 0, W, H, "FFFFFF")

    # Encabezado: píldoras de ejes + número y tiempo, con la línea azul inferior (como el HTML)
    x, y0, alto = MX, H * 0.03, 0.27
    for k, nombre in EJES:
        etiqueta = f"Eje {k} · {nombre}"
        ancho = 0.25 + len(etiqueta) * 0.062
        if "todos" in ejes or str(k) in principal:
            rel, bor, col, neg = AZUL, AZUL, "FFFFFF", True
        elif str(k) in usa:
            rel, bor, col, neg = "DBE8F6", "9CC0E6", OSCURO, False
        else:
            rel, bor, col, neg = None, "C9D2DC", "8A94A0", False
        p = caja(s, x, y0, ancho, alto, rel, bor, radio=True)
        tf = p.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        runs(tf.paragraphs[0], etiqueta, 1.6 * VH, col, neg)
        x += ancho + W * 0.005
    texto(s, W - MX - 2, y0, 2, alto, f"{num}/10 · {tiempo}", 1.6 * VH, GRIS, alin=PP_ALIGN.RIGHT, ancla=MSO_ANCHOR.MIDDLE)
    caja(s, MX, y0 + alto + 0.08, W - 2 * MX, 0.03, AZUL)

    # Pregunta (h1) y título (h2)
    texto(s, MX, 0.72, W - 2 * MX, 0.45, pregunta, 3.8 * VH, OSCURO, negrita=True)
    texto(s, MX, 1.2, W - 2 * MX, 0.32, titulo, 2.3 * VH, AZUL, negrita=True)

    # Cuerpo: viñetas (30%) + diagrama
    cuerpo_y, cuerpo_h = 1.62, 4.95
    ancho_txt = (W - 2 * MX) * 0.30
    tb = s.shapes.add_textbox(Inches(MX), Inches(cuerpo_y), Inches(ancho_txt), Inches(cuerpo_h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    for i, item in enumerate(contenido):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1 * VH)
        p.line_spacing = 1.15
        runs(p, "• " + item, 2 * VH, "1B1B1B")
    dx = MX + ancho_txt + W * 0.02
    dw = W - MX - dx
    dh = min(cuerpo_h, dw * 560 / 1000)
    dw = dh * 1000 / 560
    s.shapes.add_picture(png(DIAGRAMAS[num]), Inches(dx + (W - MX - dx - dw) / 2), Inches(cuerpo_y + (cuerpo_h - dh) / 2),
                         Inches(dw), Inches(dh))

    # Puente a la siguiente diapositiva
    p = caja(s, MX, 6.72, W - 2 * MX, 0.48, "EEF4FB")
    tf = p.text_frame
    tf.margin_left = Inches(0.15)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    runs(tf.paragraphs[0], "→ " + puente, 2 * VH, OSCURO, italica=True)

    s.notes_slide.notes_text_frame.text = guion

salida = AQUI / "presentacion-banco-plus.pptx"
prs.save(salida)
print(f"{salida.name}: {len(prs.slides)} diapositivas")
